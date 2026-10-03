"""RAG over files from WhatsApp and your computer.

@tool ingest_local_file, search_documents — plus ingest_file(), the
plain function whatsapp_server/webhook.py calls when a document/image
arrives as a WhatsApp attachment.

Extracts text from PDF, DOCX, XLSX, images (OCR via Tesseract), and
plain text/markdown; splits it into overlapping chunks; embeds and
stores each chunk in its own ChromaDB collection — separate from
tools/memory.py's short user-authored notes, since these are longer,
externally-sourced documents with their own metadata (which file, which
channel, when).
"""

import os
import uuid
from datetime import datetime

import chromadb
import openpyxl
import pytesseract
from docx import Document as DocxDocument
from langchain_core.tools import tool
from PIL import Image
from pypdf import PdfReader

from core.embeddings import VoyageEmbeddingFunction

_DB_PATH = "data/chroma"
_COLLECTION_NAME = "documents"
CHUNK_SIZE = 1000  # characters per chunk
CHUNK_OVERLAP = 150  # characters shared between consecutive chunks

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".png", ".jpg", ".jpeg"}

_collection = None


def _get_collection():
    global _collection
    if _collection is None:
        client = chromadb.PersistentClient(path=_DB_PATH)
        _collection = client.get_or_create_collection(
            name=_COLLECTION_NAME, embedding_function=VoyageEmbeddingFunction()
        )
    return _collection


def _extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        reader = PdfReader(path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == ".docx":
        doc = DocxDocument(path)
        return "\n".join(p.text for p in doc.paragraphs)
    if ext == ".xlsx":
        wb = openpyxl.load_workbook(path, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                lines.append(", ".join(str(c) for c in row if c is not None))
        return "\n".join(lines)
    if ext in {".png", ".jpg", ".jpeg"}:
        # OCR only — good for scans/screenshots with clear text. For
        # messier photos (whiteboards, handwriting), swap this for a
        # Claude vision call instead; nothing else here needs to change.
        try:
            with Image.open(path) as image:  # closed afterwards, so the file can be deleted on Windows
                return pytesseract.image_to_string(image)
        except pytesseract.TesseractNotFoundError:
            # ValueError is this module's "tell the user" error (the webhook replies with it as-is).
            raise ValueError(
                "I can't read text from images yet: the Tesseract OCR program isn't installed on the "
                "Jarvis computer (Windows installer: github.com/UB-Mannheim/tesseract/wiki)."
            ) from None
    if ext in {".txt", ".md"}:
        with open(path, encoding="utf-8", errors="ignore") as f:
            return f.read()
    raise ValueError(f"Unsupported file type: {ext} (supported: {sorted(SUPPORTED_EXTENSIONS)})")


def _chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = text.strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        chunks.append(text[start : start + chunk_size])
        start += chunk_size - overlap
    return chunks


def store_text(text: str, filename: str, source: str) -> int:
    """Chunks and stores already-extracted text directly — used by
    ingest_file() after extraction, and by whatsapp_server/webhook.py
    for voice notes (transcribed text has no file extension to dispatch
    on, so it skips _extract_text and comes straight here)."""
    chunks = _chunk_text(text)
    if not chunks:
        return 0
    ids = [str(uuid.uuid4()) for _ in chunks]
    metadatas = [
        {
            "source_file": filename,
            "source": source,
            "chunk_index": i,
            "ingested_at": datetime.now().isoformat(),
        }
        for i in range(len(chunks))
    ]
    _get_collection().add(documents=chunks, ids=ids, metadatas=metadatas)
    return len(chunks)


def ingest_file(path: str, source: str = "computer") -> int:
    """Extracts, chunks, and stores a file. Returns the number of chunks
    stored (0 if the file had no extractable text). `source` is plain
    metadata — "computer" or "whatsapp" — so search results can say
    where something came from. This is the function
    whatsapp_server/webhook.py calls directly; ingest_local_file below
    is the @tool wrapper for the agent to call the same thing on request.
    """
    text = _extract_text(path)
    return store_text(text, os.path.basename(path), source)


@tool
def ingest_local_file(path: str) -> str:
    """Read a file from the computer (PDF, DOCX, XLSX, TXT, MD, or an
    image) and make its contents searchable for future questions.

    Args:
        path: Full path to the file
    """
    if not os.path.exists(path):
        return f"File not found: {path}"
    try:
        count = ingest_file(path, source="computer")
    except ValueError as e:
        return str(e)
    if count == 0:
        return f"No text could be extracted from {os.path.basename(path)}."
    return f"Ingested {os.path.basename(path)} — {count} chunk(s) now searchable."


@tool
def search_documents(query: str, n_results: int = 5) -> str:
    """Search everything ingested from WhatsApp and the computer to help
    answer a question. Returns the most relevant passages, each labeled
    with its source file and where it came from.

    Args:
        query: What you're trying to find or answer
        n_results: How many passages to return (default 5)
    """
    results = _get_collection().query(query_texts=[query], n_results=n_results)
    documents = results.get("documents", [[]])[0]
    metadatas = results.get("metadatas", [[]])[0]
    if not documents:
        return "Nothing relevant found in ingested documents."
    return "\n\n".join(
        f"[{meta.get('source_file', '?')}, via {meta.get('source', '?')}]:\n{doc}"
        for doc, meta in zip(documents, metadatas)
    )
